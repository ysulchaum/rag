import json
from collections.abc import Mapping
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

LOG_DIRECTORY = Path(__file__).parent / "log"
LOG_FILE = LOG_DIRECTORY / "rag_query_log.xlsx"
SHEET_NAME = "Queries"
HEADERS = (
	"Timestamp (UTC)",
	"Query",
	"Needs Retrieval",
	"Context Docs (JSON)",
	"Generation",
	"Jev Model",
	"Jev Input Tokens",
	"Jev Output Tokens",
	"Jev Total Tokens",
	"LLM Provider",
	"LLM Model",
	"LLM Prompt Tokens",
	"LLM Completion Tokens",
	"LLM Total Tokens",
)


def append_query_log(
	state: Mapping[str, object], workbook_path: str | Path = LOG_FILE
) -> None:
	path = Path(workbook_path)
	path.parent.mkdir(parents=True, exist_ok=True)

	if path.exists():
		workbook = load_workbook(path)
		worksheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
	else:
		workbook = Workbook()
		worksheet = workbook.active
		worksheet.title = SHEET_NAME
		worksheet.append(HEADERS)
		for cell in worksheet[1]:
			cell.font = Font(bold=True)

	context_docs = json.dumps(state.get("context_docs", []), ensure_ascii=False)
	values = (
		datetime.now(UTC).isoformat(timespec="seconds"),
		state.get("query"),
		state.get("needs_retrieval"),
		context_docs,
		state.get("generation"),
		state.get("jev_model"),
		state.get("jev_input_tokens"),
		state.get("jev_output_tokens"),
		state.get("jev_total_tokens"),
		"DeepSeek",
		state.get("llm_model"),
		state.get("llm_prompt_tokens"),
		state.get("llm_completion_tokens"),
		state.get("llm_total_tokens"),
	)

	row_number = worksheet.max_row + 1
	for column, value in enumerate(values, start=1):
		cell = worksheet.cell(row=row_number, column=column, value=value)
		if isinstance(value, str):
			cell.data_type = "s"

	worksheet.freeze_panes = "A2"
	worksheet.auto_filter.ref = worksheet.dimensions
	workbook.save(path)
